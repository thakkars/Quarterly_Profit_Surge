import time
import traceback
import pandas as pd
import numpy as np
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import InvalidSessionIdException
from datetime import datetime
import os
import json
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Side

if os.environ.get("GITHUB_ACTIONS", "false").lower() != "true":
    from webdriver_manager.chrome import ChromeDriverManager

# ═══════════════════════════════════════════════════════════
# ENVIRONMENT DETECTION & CONFIG
# ═══════════════════════════════════════════════════════════
IS_GITHUB_ACTIONS = os.environ.get("GITHUB_ACTIONS", "false").lower() == "true"

USERNAME   = os.environ.get("SCREENER_USERNAME", "deepak.khatwani2712@gmail.com")
PASSWORD   = os.environ.get("SCREENER_PASSWORD", "gaurav23")
SCREEN_URL = "https://www.screener.in/screens/3647191/equity-research-github/"

if IS_GITHUB_ACTIONS:
    FOLDER_PATH = "/tmp/screener_data"
else:
    FOLDER_PATH = r"D:\Investment Research\Python Projects\Operating Profit more than 70% jump"

MASTER_FILE_PATH   = os.path.join(FOLDER_PATH, "Master File.xlsx")
RAW_FILE_PATH      = os.path.join(FOLDER_PATH, "Raw Data.xlsx")
RAW_JSON_PATH      = os.path.join(FOLDER_PATH, "rawdata.json")
PROGRESS_JSON_PATH = os.path.join(FOLDER_PATH, "progress.json")
INDUSTRY_FILE_PATH = os.path.join(FOLDER_PATH, "Screener_Industry_Data.xlsx")

os.makedirs(FOLDER_PATH, exist_ok=True)

HEADLESS   = IS_GITHUB_ACTIONS
PAGE_DELAY = 1.0

# ─── R2 CONFIG ─────────────────────────────────────────────
R2_BUCKET     = os.environ.get("R2_BUCKET_NAME",       "quarterly-profit-surge")
R2_ACCESS_KEY = os.environ.get("R2_ACCESS_KEY_ID",     "")
R2_SECRET_KEY = os.environ.get("R2_SECRET_ACCESS_KEY", "")
R2_ENDPOINT   = os.environ.get("R2_ENDPOINT_URL",      "")

R2_MASTER_KEY    = "Master File.xlsx"
R2_RAW_KEY       = "Raw Data.xlsx"
R2_RAW_JSON_KEY  = "rawdata.json"
R2_INDUSTRY_KEY  = "Screener_Industry_Data.xlsx"
R2_PROGRESS_KEY  = "progress.json"

# ─── MSI CONFIG (guest mode scraping) ──────────────────────
MSI_ENABLED   = True
MSI_BASE_URL  = "https://marketsmithindia.com/mstool/eval"
MSI_COLS      = [
    "MSI EPS Strength",
    "MSI Price Strength",
    "MSI Buyer Demand",
    "MSI Group Rank",
    "MSI Master Score",
]

# ═══════════════════════════════════════════════════════════
# PROGRESS TELEMETRY
# ═══════════════════════════════════════════════════════════
def push_progress(phase, completed, total, current_company="", status="running", elapsed_sec=0):
    """
    Write progress.json locally and upload to R2.
    phase: 'initialising' | 'collecting' | 'scraping' | 'enriching' | 'finalising' | 'complete' | 'failed'
    """
    pct = round((completed / total * 100)) if total > 0 else 0
    eta_sec = 0
    if completed > 0 and total > completed and elapsed_sec > 0:
        rate = elapsed_sec / completed          # seconds per stock
        eta_sec = int(rate * (total - completed))

    payload = {
        "phase":           phase,
        "status":          status,
        "completed":       completed,
        "total":           total,
        "pct":             pct,
        "current_company": current_company,
        "elapsed_sec":     int(elapsed_sec),
        "eta_sec":         eta_sec,
        "ts":              datetime.utcnow().isoformat() + "Z",
    }
    try:
        with open(PROGRESS_JSON_PATH, "w") as f:
            json.dump(payload, f)
        if IS_GITHUB_ACTIONS:
            r2_upload_progress()
    except Exception as e:
        print(f"[WARN]  Progress push failed: {e}")

def r2_upload_progress():
    """Lightweight upload of just progress.json — called after every stock."""
    try:
        import boto3
        client = boto3.client(
            "s3",
            endpoint_url          = R2_ENDPOINT,
            aws_access_key_id     = R2_ACCESS_KEY,
            aws_secret_access_key = R2_SECRET_KEY,
            region_name           = "auto",
        )
        client.upload_file(
            PROGRESS_JSON_PATH, R2_BUCKET, R2_PROGRESS_KEY,
            ExtraArgs={
                'ContentType': 'application/json',
                'CacheControl': 'no-cache, no-store, must-revalidate'
            }
        )
    except Exception as e:
        print(f"[WARN]  R2 progress upload failed: {e}")

# ═══════════════════════════════════════════════════════════
# R2 HELPERS
# ═══════════════════════════════════════════════════════════
def get_r2_client():
    import boto3
    return boto3.client(
        "s3",
        endpoint_url          = R2_ENDPOINT,
        aws_access_key_id     = R2_ACCESS_KEY,
        aws_secret_access_key = R2_SECRET_KEY,
        region_name           = "auto",
    )

def r2_download(r2_key, local_path):
    try:
        client = get_r2_client()
        client.download_file(R2_BUCKET, r2_key, local_path)
        print(f"[DOWN]  Downloaded from R2: {r2_key} -> {local_path}")
        return True
    except Exception as e:
        print(f"[WARN]  R2 download failed for {r2_key}: {e}")
        return False

def r2_upload(local_path, r2_key):
    try:
        if not os.path.exists(local_path):
            print(f"[WARN]  Skipping R2 upload - file not found: {local_path}")
            return False
        client = get_r2_client()
        extra_args = {}
        if r2_key.endswith('.json'):
            extra_args = {
                'ContentType': 'application/json',
                'CacheControl': 'no-cache, no-store, must-revalidate'
            }
        client.upload_file(local_path, R2_BUCKET, r2_key, ExtraArgs=extra_args)
        print(f"[UP]  Uploaded to R2 : {local_path} -> {r2_key}")
        return True
    except Exception as e:
        print(f"[WARN]  R2 upload failed for {r2_key}: {e}")
        return False

if IS_GITHUB_ACTIONS:
    print("\n[DOWN] Downloading existing files from Cloudflare R2...")
    r2_download(R2_MASTER_KEY,   MASTER_FILE_PATH)
    r2_download(R2_RAW_KEY,      RAW_FILE_PATH)
    r2_download(R2_INDUSTRY_KEY, INDUSTRY_FILE_PATH)

# ═══════════════════════════════════════════════════════════
# MAPPING & ENRICHMENT
# ═══════════════════════════════════════════════════════════
def _get_industry_df():
    file_path = INDUSTRY_FILE_PATH
    if not os.path.exists(file_path):
        dir_name = os.path.dirname(file_path)
        fallbacks = [
            os.path.join(dir_name, "screener_industry.xlsx"),
            os.path.join(dir_name, "Screener_Industry_Data.csv"),
            os.path.join(dir_name, "Screener_Industry_Data.xlsx")
        ]
        for f in fallbacks:
            if os.path.exists(f):
                file_path = f
                break
        else:
            return None, ""
    try:
        if file_path.endswith('.csv'):
            df = pd.read_csv(file_path)
        else:
            xl = pd.ExcelFile(file_path)
            sheet = "Screener Industry" if "Screener Industry" in xl.sheet_names else xl.sheet_names[0]
            df = pd.read_excel(xl, sheet_name=sheet)
        return df, file_path
    except Exception as e:
        print(f"[WARN] Error loading industry data file {file_path}: {e}")
        return None, ""

def load_industry_maps():
    df, path = _get_industry_df()
    if df is None:
        print("[WARN] No industry data file found - industry maps empty")
        return {}, {}
    try:
        # Standardize column mapping
        name_col = None
        for c in df.columns:
            c_clean = str(c).strip().lower()
            if c_clean in ('full name', 'company', 'name'):
                name_col = c
                break
        if not name_col:
            return {}, {}
            
        df["Name_Clean"] = df[name_col].astype(str).str.strip().str.lower()
        
        industry_map = {}
        group_map = {}
        for _, row in df.iterrows():
            k = str(row["Name_Clean"]).strip()
            if not k or k == "nan":
                continue
            if "Industry" in row:
                industry_map[k] = row["Industry"]
            if "Industry Group" in row:
                group_map[k] = row["Industry Group"]
                
        # Also map short Name if Company and Name columns exist
        if 'Company' in df.columns and 'Name' in df.columns:
            for _, row in df.iterrows():
                short_name = str(row['Name']).strip().lower()
                if short_name and short_name != 'nan':
                    if "Industry" in row:
                        industry_map[short_name] = row["Industry"]
                    if "Industry Group" in row:
                        group_map[short_name] = row["Industry Group"]
                        
        print(f"[OK] Loaded industry mappings from {os.path.basename(path)}")
        return industry_map, group_map
    except Exception as e:
        print(f"[WARN] Could not parse industry maps: {e}")
        return {}, {}

def load_symbol_maps():
    df, path = _get_industry_df()
    if df is None:
        print("[WARN] No symbol mapping file found - symbol maps empty")
        return {}
    try:
        # Standardize column mapping
        col_mapping = {}
        for c in df.columns:
            c_clean = str(c).strip().lower()
            if c_clean in ('full name', 'company', 'name'):
                col_mapping['name'] = c
            elif c_clean in ('nse code', 'nse_code'):
                col_mapping['nse'] = c
            elif c_clean in ('bse code', 'bse_code'):
                col_mapping['bse'] = c

        name_col = col_mapping.get('name')
        nse_col = col_mapping.get('nse')
        bse_col = col_mapping.get('bse')

        if not name_col:
            print("[WARN] Name column not found in mapping file")
            return {}

        mapping = {}
        for _, row in df.iterrows():
            name_val = str(row[name_col]).strip()
            if not name_val or name_val.lower() == "nan":
                continue

            nse = str(row[nse_col]).strip() if nse_col and pd.notna(row[nse_col]) else ""
            bse = str(row[bse_col]).strip() if bse_col and pd.notna(row[bse_col]) else ""

            if nse.lower() in ("nan", "none", "0", ""):
                nse = ""
            if bse.lower() in ("nan", "none", "0", ""):
                bse = ""
            if bse and bse.endswith(".0"):
                bse = bse[:-2]

            screener_slug = nse if nse else bse

            mapping[name_val.lower()] = {
                "NSE_Code":      nse,
                "BSE_Code":      bse,
                "Screener_Name": screener_slug
            }

        # Also support short Name mapping if Company column is present
        if 'Company' in df.columns and 'Name' in df.columns:
            for _, row in df.iterrows():
                short_name = str(row['Name']).strip()
                if short_name and short_name.lower() != 'nan':
                    nse = str(row[nse_col]).strip() if nse_col and pd.notna(row[nse_col]) else ""
                    bse = str(row[bse_col]).strip() if bse_col and pd.notna(row[bse_col]) else ""
                    if nse.lower() in ("nan", "none", "0", ""): nse = ""
                    if bse.lower() in ("nan", "none", "0", ""): bse = ""
                    if bse and bse.endswith(".0"): bse = bse[:-2]
                    screener_slug = nse if nse else bse
                    
                    mapping[short_name.lower()] = {
                        "NSE_Code":      nse,
                        "BSE_Code":      bse,
                        "Screener_Name": screener_slug
                    }

        print(f"[OK] Loaded {len(mapping)} symbol mappings from {os.path.basename(path)}")
        return mapping
    except Exception as e:
        print(f"[WARN] Could not parse symbol maps: {e}")
        return {}

def enrich_with_data(row_dict, industry_map, group_map, symbol_map):
    full_name_clean  = str(row_dict.get("Company (Page)", "")).strip().lower()
    short_name_clean = str(row_dict.get("Name", "")).strip().lower()
    row_dict["Industry"]       = industry_map.get(full_name_clean, "Not Found")
    row_dict["Industry Group"] = group_map.get(full_name_clean, "Not Found")
    mapped = symbol_map.get(full_name_clean) or symbol_map.get(short_name_clean) or {}
    row_dict["NSE Code Exact"]      = mapped.get("NSE_Code", "")
    row_dict["Screener Name Exact"] = mapped.get("Screener_Name", "")
    if not row_dict["Screener Name Exact"]:
        row_dict["Screener Name Exact"] = str(row_dict.get("Name", "")).strip()
    return row_dict

# ═══════════════════════════════════════════════════════════
# SELENIUM DRIVER INIT
# ═══════════════════════════════════════════════════════════
def init_driver():
    opts = Options()
    if HEADLESS:
        opts.add_argument("--headless=new")
    opts.add_argument("--disable-gpu")
    opts.add_argument("--window-size=1600,900")
    opts.add_argument("--no-sandbox")
    opts.add_argument("--disable-dev-shm-usage")
    if IS_GITHUB_ACTIONS:
        return webdriver.Chrome(options=opts)
    else:
        return webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=opts)

driver = init_driver()
wait   = WebDriverWait(driver, 15)

# ═══════════════════════════════════════════════════════════
# MARKETSMITH INDIA (MSI) GUEST SCRAPER
# ═══════════════════════════════════════════════════════════
def msi_get_metric_value(msi_driver, msi_wait, header_text):
    try:
        h4 = msi_wait.until(
            EC.presence_of_element_located(
                (By.XPATH, f"//h4[normalize-space()='{header_text}']")
            )
        )
        h3 = h4.find_element(By.XPATH, "following::h3[1]")
        spans = h3.find_elements(By.TAG_NAME, "span")
        if len(spans) >= 2:
            return spans[1].text.strip()
        elif spans:
            return spans[0].text.split(":")[-1].strip()
    except Exception:
        pass
    return ""

def msi_get_master_score(msi_driver):
    try:
        ms = msi_driver.find_element(By.CSS_SELECTOR, "#masterScoreIndicator .masters b")
        return ms.text.strip()
    except Exception:
        pass
    try:
        h2 = msi_driver.find_element(By.CSS_SELECTOR, "h2.contentMasters b")
        parts = h2.text.strip().split()
        if len(parts) >= 3:
            return parts[2].rstrip(":")
        return ""
    except Exception:
        return ""

def msi_scrape_symbol(msi_driver, msi_wait, symbol):
    result = {
        "MSI EPS Strength":   "",
        "MSI Price Strength": "",
        "MSI Buyer Demand":   "",
        "MSI Group Rank":     "",
        "MSI Master Score":   "",
    }
    if not symbol:
        return result
    url = f"{MSI_BASE_URL}/{symbol.lower()}/evaluation.jsp"
    try:
        msi_driver.get(url)
        msi_wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, "span")))
        time.sleep(1.5)
        result["MSI EPS Strength"]   = msi_get_metric_value(msi_driver, msi_wait, "EPS Strength")
        result["MSI Price Strength"] = msi_get_metric_value(msi_driver, msi_wait, "Price Strength")
        result["MSI Buyer Demand"]   = msi_get_metric_value(msi_driver, msi_wait, "Buyer Demand")
        result["MSI Group Rank"]     = msi_get_metric_value(msi_driver, msi_wait, "Group Rank")
        result["MSI Master Score"]   = msi_get_master_score(msi_driver)
    except Exception as e:
        print(f"  MSI error for {symbol}: {e}")
    return result

def attach_msi_metrics(row_dict, msi_cache, msi_driver, msi_wait):
    if not MSI_ENABLED or msi_driver is None:
        for c in MSI_COLS:
            row_dict.setdefault(c, "")
        return row_dict
    symbol = str(
        row_dict.get("NSE Code Exact") or row_dict.get("Screener Name Exact") or ""
    ).strip().upper()
    if not symbol:
        for c in MSI_COLS:
            row_dict.setdefault(c, "")
        return row_dict
    if symbol in msi_cache:
        row_dict.update(msi_cache[symbol])
        return row_dict
    metrics = msi_scrape_symbol(msi_driver, msi_wait, symbol)
    msi_cache[symbol] = metrics
    row_dict.update(metrics)
    return row_dict

# ═══════════════════════════════════════════════════════════
# SCRAPER LOGIC (SCREENER)
# ═══════════════════════════════════════════════════════════
def login():
    driver.get("https://www.screener.in/login/")
    try:
        wait.until(EC.presence_of_element_located((By.NAME, "username"))).send_keys(USERNAME)
        driver.find_element(By.NAME, "password").send_keys(PASSWORD)
        driver.find_element(By.NAME, "password").submit()
        time.sleep(3)
    except:
        pass

def safe_get(url):
    global driver, wait
    try:
        driver.get(url)
        return True
    except InvalidSessionIdException:
        try:
            driver.quit()
        except:
            pass
        driver = init_driver()
        wait   = WebDriverWait(driver, 15)
        login()
        return safe_get(url)
    except:
        return False

def get_company_full_name():
    selectors = ["h1.margin-0", "h1.shrink-text", "h1", ".company-name h1"]
    for sel in selectors:
        try:
            el   = wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, sel)))
            name = el.text.strip()
            if name and name.lower() != "unknown":
                return name
        except:
            continue
    return "Unknown"

def get_top_dashboard_metrics():
    metrics = {"Promoter Holding": 0.0, "Pledged %": 0.0, "WC Days": 0.0}
    try:
        ratio_items = driver.find_elements(By.CSS_SELECTOR, "ul#top-ratios li")
        for li in ratio_items:
            try:
                label  = li.find_element(By.CSS_SELECTOR, "span.name").text.lower().strip()
                val_el = li.find_element(By.CSS_SELECTOR, "span.number").text.strip().replace(",", "")
                value  = float(val_el) if val_el else 0.0
                if "promoter holding" in label:
                    metrics["Promoter Holding"] = value
                elif "pledged" in label:
                    metrics["Pledged %"] = value
                elif "wc days" in label:
                    metrics["WC Days"] = value
            except:
                continue
    except:
        pass
    return metrics

def extract_operating_profit():
    try:
        table = wait.until(EC.presence_of_element_located((By.XPATH, "//section[@id='quarters']//table")))
    except:
        return {}
    headers    = [th.text.strip() for th in table.find_elements(By.TAG_NAME, "th")]
    rows       = table.find_elements(By.TAG_NAME, "tr")
    target_row = next((r for r in rows if "operating profit" in r.text.lower()), None)
    if not target_row:
        target_row = next((r for r in rows if "financing profit" in r.text.lower()), None)
    if not target_row:
        return {}
    values = [c.text.strip() for c in target_row.find_elements(By.TAG_NAME, "td")]
    return {h: v for h, v in zip(headers[1:], values[1:])}

def collect_all_company_data():
    if not safe_get(SCREEN_URL):
        return []
    company_list, seen_urls = [], set()
    while True:
        for a in driver.find_elements(By.CSS_SELECTOR, "a[href*='/company/']"):
            href, short_name = a.get_attribute("href"), a.text.strip()
            if href and "/company/" in href and href not in seen_urls:
                seen_urls.add(href)
                company_list.append({"short_name": short_name, "url": href})
        try:
            next_btn = driver.find_element(By.XPATH, "//a[contains(text(), 'Next')]")
            if "disabled" in (next_btn.get_attribute("class") or "").lower():
                break
            driver.execute_script("arguments[0].click();", next_btn)
            time.sleep(1.5)
        except:
            break
    return company_list

def scrape_company(company_info):
    url, short_name = company_info['url'], company_info['short_name']
    if not safe_get(url):
        return {
            "Name": short_name,
            "Company (Page)": "Load Failed",
            "date": datetime.now().strftime("%Y-%m-%d"),
        }
    try:
        time.sleep(1.5)
        try:
            driver.find_element(By.LINK_TEXT, "Consolidated").click()
            time.sleep(1.5)
        except:
            try:
                driver.find_element(By.PARTIAL_LINK_TEXT, "Consolidated").click()
                time.sleep(1.5)
            except:
                pass
        full_name = get_company_full_name()
        if full_name == "Unknown":
            slug = url.rstrip("/").split("/")[-1]
            full_name = slug if slug else short_name
        return {
            "Name": short_name,
            "Company (Page)": full_name,
            "Industry Group": "",
            "Industry": "",
            **extract_operating_profit(),
            **get_top_dashboard_metrics(),
            "date": datetime.now().strftime("%Y-%m-%d"),
        }
    except:
        return {
            "Name": short_name,
            "Company (Page)": short_name,
            "date": datetime.now().strftime("%Y-%m-%d"),
        }

def detect_latest_quarter(row_dict):
    non_quarter_keys = {
        "Sr no", "Name", "Company (Page)", "Industry Group", "Industry",
        "Promoter Holding", "Pledged %", "WC Days", "date",
        "NSE Code Exact", "Screener Name Exact",
        *MSI_COLS,
    }
    quarter_candidates = []
    for k in row_dict:
        if k in non_quarter_keys or "(>70%)" in k:
            continue
        try:
            quarter_candidates.append((pd.to_datetime(k, format="%b %Y"), k))
        except:
            pass
    if not quarter_candidates:
        return None
    quarter_candidates.sort(key=lambda x: x[0], reverse=True)
    return quarter_candidates[0][0].strftime("%b %y")

def get_existing_companies_in_sheet(sheet_name):
    if not os.path.exists(RAW_FILE_PATH):
        return set()
    try:
        xl = pd.ExcelFile(RAW_FILE_PATH)
        if sheet_name not in xl.sheet_names:
            return set()
        df = pd.read_excel(RAW_FILE_PATH, sheet_name=sheet_name)
        if "Name" in df.columns:
            return set(df["Name"].dropna().astype(str).str.strip())
        return set()
    except:
        return set()

def transform_and_order_data(df):
    base_cols   = [
        "Sr no", "date", "Name", "Company (Page)", "Industry", "Industry Group",
        "NSE Code Exact", "Screener Name Exact",
    ]
    metric_cols = [
        "Pledged %",
        "Promoter Holding",
        "WC Days",
        *MSI_COLS,
    ]
    for col in base_cols + metric_cols:
        if col not in df.columns:
            df[col] = ""
    quarter_cols_raw = [
        c for c in df.columns
        if c not in base_cols
        and c not in metric_cols
        and "(>70%)" not in c
        and c != "70% Surge"
    ]
    parsed_dates = []
    for c in quarter_cols_raw:
        try:
            parsed_dates.append((pd.to_datetime(c, format="%b %Y"), c))
        except:
            pass
    parsed_dates.sort(key=lambda x: x[0], reverse=True)
    sorted_quarter_cols = [x[1] for x in parsed_dates]
    final_cols = base_cols + sorted_quarter_cols + metric_cols
    df = df[[c for c in final_cols if c in df.columns]].copy()
    df["Sr no"] = range(1, len(df) + 1)
    return df

def append_rows_to_raw_excel(rows, sheet_name):
    if not rows:
        return
    df_new = transform_and_order_data(pd.DataFrame(rows))
    if os.path.exists(RAW_FILE_PATH):
        try:
            wb = load_workbook(RAW_FILE_PATH)
            if sheet_name in wb.sheetnames:
                df_existing = pd.read_excel(RAW_FILE_PATH, sheet_name=sheet_name)
                df_final    = transform_and_order_data(
                    pd.concat([df_existing, df_new], ignore_index=True)
                )
            else:
                df_final = df_new
            wb.close()
            with pd.ExcelWriter(
                RAW_FILE_PATH,
                engine='openpyxl',
                mode='a',
                if_sheet_exists='replace'
            ) as writer:
                df_final.to_excel(writer, sheet_name=sheet_name, index=False)
        except:
            return
    else:
        with pd.ExcelWriter(RAW_FILE_PATH, engine='openpyxl') as writer:
            df_new.to_excel(writer, sheet_name=sheet_name, index=False)
    wb = load_workbook(RAW_FILE_PATH)
    ws = wb[sheet_name]
    for row in ws.iter_rows():
        for cell in row:
            cell.border = Border(
                left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'),  bottom=Side(style='thin')
            )
            if cell.row == 1 or cell.col_idx == 1:
                cell.alignment = Alignment(horizontal='center', vertical='center')
    wb.save(RAW_FILE_PATH)

def generate_raw_json(sheet_name):
    if not os.path.exists(RAW_FILE_PATH):
        return
    try:
        xl = pd.ExcelFile(RAW_FILE_PATH)
        if sheet_name not in xl.sheet_names:
            sheet_name = xl.sheet_names[-1]
        df = pd.read_excel(RAW_FILE_PATH, sheet_name=sheet_name)
        df.columns = df.columns.astype(str).str.strip()
        df = df.fillna(0)
        df.to_json(RAW_JSON_PATH, orient="records")
        print(f"[OK] Saved rawdata.json for Dashboard (Sheet: {sheet_name})")
    except Exception as e:
        print(f"[WARN]  Error saving JSON: {e}")

# ═══════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════
run_start = time.time()
try:
    push_progress("initialising", 0, 1, "Authenticating with Screener.in...")

    INDUSTRY_MAP, GROUP_MAP = load_industry_maps()
    SYMBOL_MAP              = load_symbol_maps()

    msi_driver = init_driver() if MSI_ENABLED else None
    msi_wait   = WebDriverWait(msi_driver, 20) if MSI_ENABLED and msi_driver else None
    msi_cache  = {}

    login()

    push_progress("collecting", 0, 1, "Enumerating universe from Screener screen...")
    company_data_list = collect_all_company_data()
    total = len(company_data_list)
    print(f"[STATS] Universe identified: {total} equities")
    push_progress("collecting", 0, total, f"Universe identified: {total} equities")

    if company_data_list:
        push_progress("scraping", 0, total, company_data_list[0]['short_name'])

        first_scraped = scrape_company(company_data_list[0])
        first_row     = enrich_with_data(first_scraped, INDUSTRY_MAP, GROUP_MAP, SYMBOL_MAP)
        first_row     = attach_msi_metrics(first_row, msi_cache, msi_driver, msi_wait)
        push_progress("scraping", 1, total, first_row.get("Company (Page)", ""), elapsed_sec=time.time()-run_start)

        target_sheet   = detect_latest_quarter(first_row) or datetime.now().strftime("%b %y")
        existing_names = get_existing_companies_in_sheet(target_sheet)
        pending_rows   = [] if str(first_row.get("Name", "")).strip() in existing_names else [first_row]

        remaining = [
            c for c in company_data_list[1:]
            if str(c['short_name']).strip() not in existing_names
        ]

        completed_count = 1
        for i, info in enumerate(remaining, start=2):
            scraped  = scrape_company(info)
            enriched = enrich_with_data(scraped, INDUSTRY_MAP, GROUP_MAP, SYMBOL_MAP)
            enriched = attach_msi_metrics(enriched, msi_cache, msi_driver, msi_wait)
            pending_rows.append(enriched)
            completed_count += 1

            push_progress(
                "scraping",
                completed_count,
                total,
                enriched.get("Company (Page)", info['short_name']),
                elapsed_sec=time.time() - run_start
            )

            if len(pending_rows) >= 10:
                append_rows_to_raw_excel(pending_rows, target_sheet)
                pending_rows = []

        if pending_rows:
            append_rows_to_raw_excel(pending_rows, target_sheet)

        push_progress("finalising", total, total, "Serialising dataset & committing to R2...")
        generate_raw_json(target_sheet)

        if IS_GITHUB_ACTIONS:
            r2_upload(RAW_FILE_PATH,    R2_RAW_KEY)
            r2_upload(RAW_JSON_PATH,    R2_RAW_JSON_KEY)
            r2_upload(MASTER_FILE_PATH, R2_MASTER_KEY)

        push_progress("complete", total, total, "Dataset refreshed successfully.", status="complete", elapsed_sec=time.time()-run_start)

except Exception as e:
    print(f"\n[ERROR] Fatal error: {e}")
    traceback.print_exc()
    push_progress("failed", 0, 0, str(e), status="failed")
finally:
    try:
        driver.quit()
    except:
        pass
    try:
        if MSI_ENABLED and msi_driver:
            msi_driver.quit()
    except:
        pass
