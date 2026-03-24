import time
import random
import traceback
import pandas as pd
import numpy as np
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException, InvalidSessionIdException
from datetime import datetime
import os
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Side

# ═══════════════════════════════════════════════════════════
#  CONFIG  — credentials now come from environment variables
# ═══════════════════════════════════════════════════════════
USERNAME  = os.environ.get("SCREENER_USERNAME", "")
PASSWORD  = os.environ.get("SCREENER_PASSWORD", "")

SCREEN_URL = "https://www.screener.in/screens/742984/top-growers/"

# ── HEADLESS is always True when running on GitHub Actions cloud
IS_GITHUB_ACTIONS = os.environ.get("GITHUB_ACTIONS", "false").lower() == "true"
HEADLESS   = IS_GITHUB_ACTIONS   # True on cloud, False on your laptop
PAGE_DELAY = 1.0

# ── Paths:
if IS_GITHUB_ACTIONS:
    FOLDER_PATH = "/tmp/screener_data"
else:
    FOLDER_PATH = r"D:\Investment Research\Python Projects\Operating Profit more than 70% jump"

MASTER_FILE_PATH = os.path.join(FOLDER_PATH, "Master File.xlsx")
RAW_FILE_PATH    = os.path.join(FOLDER_PATH, "Raw Data.xlsx")

os.makedirs(FOLDER_PATH, exist_ok=True)


# ═══════════════════════════════════════════════════════════
#  DRIVER SETUP
# ═══════════════════════════════════════════════════════════
def init_driver():
    opts = Options()
    if HEADLESS:
        opts.add_argument("--headless=new")
    opts.add_argument("--disable-gpu")
    opts.add_argument("--window-size=1600,900")
    opts.add_argument("--no-sandbox")                    
    opts.add_argument("--disable-dev-shm-usage")         
    opts.add_experimental_option("excludeSwitches", ["enable-automation"])
    opts.add_experimental_option('useAutomationExtension', False)

    if IS_GITHUB_ACTIONS:
        opts.binary_location = "/usr/bin/google-chrome"
        return webdriver.Chrome(options=opts)
    else:
        from webdriver_manager.chrome import ChromeDriverManager
        return webdriver.Chrome(
            service=Service(ChromeDriverManager().install()),
            options=opts
        )

driver = init_driver()
wait   = WebDriverWait(driver, 15)


# ═══════════════════════════════════════════════════════════
#  HELPER FUNCTIONS
# ═══════════════════════════════════════════════════════════
def login():
    driver.get("https://www.screener.in/login/")
    try:
        wait.until(EC.presence_of_element_located((By.NAME, "username")))
        driver.find_element(By.NAME, "username").send_keys(USERNAME)
        driver.find_element(By.NAME, "password").send_keys(PASSWORD)
        driver.find_element(By.NAME, "password").submit()
        time.sleep(3)
    except:
        pass

def safe_text(el):
    try:
        return el.text.strip()
    except:
        return ""

def safe_get(url, retries=1):
    global driver, wait
    try:
        driver.get(url)
        return True
    except InvalidSessionIdException:
        print("⚠️ Session died. Restarting Chrome...")
        try:
            driver.quit()
        except:
            pass
        driver = init_driver()
        wait   = WebDriverWait(driver, 15)
        login()
        if retries > 0:
            return safe_get(url, retries - 1)
        return False
    except Exception as e:
        print(f"⚠️ Unexpected error in safe_get: {e}")
        return False

def get_company_full_name():
    try:
        h1 = wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, "h1.margin-0")))
        return h1.text.strip()
    except:
        return "Unknown"

def select_consolidated_if_exists():
    try:
        main_window = driver.current_window_handle
        tab = driver.find_element(By.LINK_TEXT, "Consolidated")
        tab.click()
        time.sleep(1.5)
        all_windows = driver.window_handles
        for w in all_windows:
            if w != main_window:
                driver.switch_to.window(w)
                return main_window
        return main_window
    except:
        return driver.current_window_handle


# ═══════════════════════════════════════════════════════════
#  METRIC EXTRACTION
# ═══════════════════════════════════════════════════════════
def get_top_dashboard_metrics():
    metrics = {"Promoter Holding": 0.0, "Pledged %": 0.0, "WC Days": 0.0}
    try:
        ratio_items = driver.find_elements(By.CSS_SELECTOR, "ul#top-ratios li")
        for li in ratio_items:
            try:
                name_el = li.find_element(By.CSS_SELECTOR, "span.name")
                val_el  = li.find_element(By.CSS_SELECTOR, "span.number")
                label   = name_el.text.lower().strip()
                value   = float(val_el.text.strip().replace(",", "")) if val_el.text.strip() else 0.0
                if "promoter holding" in label:
                    metrics["Promoter Holding"] = value
                elif "pledged percentage" in label or "pledged" in label:
                    metrics["Pledged %"] = value
                elif "wc days" in label or "working capital days" in label:
                    metrics["WC Days"] = value
            except:
                continue
        return metrics
    except:
        return metrics

def extract_operating_profit():
    try:
        table = wait.until(EC.presence_of_element_located(
            (By.XPATH, "//section[@id='quarters']//table")))
    except TimeoutException:
        return {}

    headers    = [th.text.strip() for th in table.find_elements(By.TAG_NAME, "th")]
    rows       = table.find_elements(By.TAG_NAME, "tr")
    target_row = None

    for r in rows:
        if "operating profit" in r.text.lower():
            target_row = r
            break
    if not target_row:
        for r in rows:
            if "financing profit" in r.text.lower():
                target_row = r
                break
    if not target_row:
        return {}

    values = [safe_text(c) for c in target_row.find_elements(By.TAG_NAME, "td")]
    return {h: v for h, v in zip(headers[1:], values[1:])}

def collect_all_company_data():
    if not safe_get(SCREEN_URL):
        return []
    wait.until(EC.presence_of_all_elements_located((By.CSS_SELECTOR, "a[href*='/company/']")))
    company_list = []
    seen_urls    = set()

    while True:
        rows = driver.find_elements(By.CSS_SELECTOR, "a[href*='/company/']")
        for a in rows:
            href       = a.get_attribute("href")
            short_name = a.text.strip()
            if href and "/company/" in href and href not in seen_urls:
                seen_urls.add(href)
                company_list.append({"short_name": short_name, "url": href})
        try:
            next_btn = driver.find_element(By.XPATH, "//a[contains(text(), 'Next')]")
            if "disabled" in (next_btn.get_attribute("class") or "").lower():
                break
            driver.execute_script("arguments[0].click();", next_btn)
            wait.until(EC.presence_of_all_elements_located(
                (By.CSS_SELECTOR, "a[href*='/company/']")))
            time.sleep(1.5)
        except:
            break

    print(f"✅ Collected {len(company_list)} companies.")
    return company_list

def scrape_company(company_info):
    url        = company_info['url']
    short_name = company_info['short_name']

    if not safe_get(url):
        return {"Name": short_name, "Company (Page)": "Load Failed"}

    main_window = driver.current_window_handle
    try:
        _           = select_consolidated_if_exists()
        full_name   = get_company_full_name()
        op_data     = extract_operating_profit()
        top_metrics = get_top_dashboard_metrics()

        all_windows = driver.window_handles
        for w in all_windows:
            if w != main_window:
                driver.close()
        driver.switch_to.window(main_window)

        return {
            "Name"           : short_name,
            "Company (Page)" : full_name,
            "Industry Group" : "",
            "Industry"       : "",
            **op_data,
            **top_metrics
        }
    except Exception as e:
        print(f"⚠️ Error scraping {short_name}: {e}")
        all_windows = driver.window_handles
        for w in all_windows:
            if w != main_window:
                driver.close()
        driver.switch_to.window(main_window)
        return {"Name": short_name, "Company (Page)": "Error"}


# ═══════════════════════════════════════════════════════════
#  EXCEL HELPERS
# ═══════════════════════════════════════════════════════════
def open_or_create_excel(path):
    if os.path.exists(path):
        return load_workbook(path)
    else:
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = "_init"
        wb.save(path)
        return wb

THIN_BORDER = Border(
    left   = Side(style='thin'),
    right  = Side(style='thin'),
    top    = Side(style='thin'),
    bottom = Side(style='thin')
)
CENTER_ALIGN = Alignment(horizontal='center', vertical='center')

def format_worksheet(ws, header_row=1):
    for row in ws.iter_rows():
        for cell in row:
            cell.border = THIN_BORDER
            if cell.row == header_row or cell.col_idx == 1:
                cell.alignment = CENTER_ALIGN

def clean_and_convert(df_in):
    def to_num(col):
        return pd.to_numeric(col, errors='coerce').fillna(0)
    for col in df_in.columns:
        col_str    = str(col).lower()
        is_quarter = (
            any(m in col_str for m in ['sep', 'jun', 'mar', 'dec'])
            and any(c.isdigit() for c in col_str)
        )
        is_metric  = col_str in ['promoter holding', 'pledged %', 'wc days', 'sr no.']
        if is_quarter or is_metric:
            df_in[col] = to_num(df_in[col])
    return df_in


# ═══════════════════════════════════════════════════════════
#  CLOUDFLARE R2 UPLOAD
# ═══════════════════════════════════════════════════════════
def upload_to_cloudflare_r2():
    access_key    = os.environ.get("R2_ACCESS_KEY_ID", "")
    secret_key    = os.environ.get("R2_SECRET_ACCESS_KEY", "")
    endpoint_url  = os.environ.get("R2_ENDPOINT_URL", "")
    bucket_name   = os.environ.get("R2_BUCKET_NAME", "quarterly-profit-surge")

    if not all([access_key, secret_key, endpoint_url]):
        print("⚠️ R2 credentials not set — skipping upload.")
        return

    try:
        import boto3
        from botocore.config import Config

        s3 = boto3.client(
            "s3",
            endpoint_url          = endpoint_url,
            aws_access_key_id     = access_key,
            aws_secret_access_key = secret_key,
            config                = Config(signature_version="s3v4"),
            region_name           = "auto"
        )

        files_to_upload = [
            (MASTER_FILE_PATH, "Master File.xlsx"),
            (RAW_FILE_PATH,    "Raw Data.xlsx"),
        ]

        for local_path, filename in files_to_upload:
            if not os.path.exists(local_path):
                print(f"⚠️ {filename} not found, skipping.")
                continue

            s3.upload_file(
                local_path,
                bucket_name,
                filename,
                ExtraArgs={"ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
            )
            print(f"✅ Uploaded {filename} to Cloudflare R2.")

    except Exception as e:
        print(f"⚠️ R2 upload failed: {e}")
        traceback.print_exc()

def upload_to_google_drive():
    upload_to_cloudflare_r2()


# ═══════════════════════════════════════════════════════════
#  MAIN
# ═══════════════════════════════════════════════════════════
try:
    login()
    company_data_list = collect_all_company_data()

    if company_data_list:
        data  = []
        total = len(company_data_list)
        for i, info in enumerate(company_data_list, start=1):
            row = scrape_company(info)
            data.append(row)
            print(f"✅ [{i}/{total}] {info['short_name']}")
            time.sleep(PAGE_DELAY + random.uniform(0.3, 0.7))

        df = pd.DataFrame(data)

        # ── 1. IDENTIFY QUARTERS & SORT ──────────────────────────────────
        non_quarter_cols = [
            "Name", "Company (Page)", "Promoter Holding",
            "Pledged %", "WC Days", "Industry Group", "Industry"
        ]
        raw_quarters  = [c for c in df.columns if c not in non_quarter_cols]
        quarter_dates = pd.to_datetime(raw_quarters, format="%b %Y", errors='coerce')
        valid_quarters = [
            (d, c) for d, c in zip(quarter_dates, raw_quarters) if not pd.isna(d)
        ]
        valid_quarters.sort(key=lambda x: x[0], reverse=True)
        sorted_quarters = [x[1] for x in valid_quarters]

        if sorted_quarters:
            latest_quarter_label = sorted_quarters[0]
        else:
            latest_quarter_label = datetime.now().strftime("%b %Y")

        print(f"\n📅 Latest quarter detected: {latest_quarter_label}")

        final_quarters = sorted_quarters[:13]
        while len(final_quarters) < 13:
            final_quarters.append(f"Empty_Q{len(final_quarters)+1}")
            df[final_quarters[-1]] = 0

        numeric_quarters = df[sorted_quarters[:13]].apply(pd.to_numeric, errors="coerce")
        df[final_quarters] = numeric_quarters

        # ── 2. CALCULATE OP PROFIT LOGIC ─────────────────────────────────
        latest_col = final_quarters[0]
        try:
            latest_dt  = pd.to_datetime(latest_col, format="%b %Y")
            target_dt  = latest_dt.replace(year=latest_dt.year - 1)
            prev_year_col = next(
                (c for c in final_quarters
                 if pd.to_datetime(c, format="%b %Y", errors='coerce') == target_dt),
                None
            )
            latest_vals = numeric_quarters[latest_col].fillna(0)
            if prev_year_col:
                prev_vals = numeric_quarters[prev_year_col].fillna(0)
                df["OP Profit > 70%"] = latest_vals > 1.7 * prev_vals
            else:
                df["OP Profit > 70%"] = False

            max_past = numeric_quarters.max(axis=1).fillna(0)
            df["Current Profit > in past 3 yrs"] = latest_vals >= max_past
        except:
            df["OP Profit > 70%"]                = False
            df["Current Profit > in past 3 yrs"] = False

        # ── 3. YoY COMPARISONS ───────────────────────────────────────────
        comparison_cols = []
        comp_pairs = [(1,5),(2,6),(3,7),(4,8),(5,9),(6,10),(7,11),(8,12)]
        for n_idx, o_idx in comp_pairs:
            n, o = n_idx - 1, o_idx - 1
            if n < len(final_quarters) and o < len(final_quarters):
                c_new, c_old = final_quarters[n], final_quarters[o]
                c_name = f"{c_new} > 1.7*{c_old}"
                if "Empty" not in c_new:
                    df[c_name] = numeric_quarters[c_new] > 1.7 * numeric_quarters[c_old]
                else:
                    df[c_name] = False
                comparison_cols.append(c_name)
            else:
                c_name = f"Comp_{len(comparison_cols)+1}"
                df[c_name] = False
                comparison_cols.append(c_name)
        final_comparisons = comparison_cols[:8]

        # ── 4. EXPORT STRUCTURE ──────────────────────────────────────────
        desired_order = (
            ["Name", "Company (Page)"]
            + final_quarters
            + ["OP Profit > 70%", "Current Profit > in past 3 yrs"]
            + final_comparisons
            + ["Promoter Holding", "Pledged %", "WC Days", "Industry Group", "Industry"]
        )
        df_export = df.reindex(columns=desired_order)

        # ── 5. ENRICHMENT (VLOOKUP & EXCEPTION) ──────────────────────────
        print("\n🔄 Performing VLOOKUP & Exception Check...")
        try:
            if os.path.exists(MASTER_FILE_PATH):
                df_source = pd.read_excel(MASTER_FILE_PATH, sheet_name="screener_all_stocks")
                if "Name" in df_source.columns and "Industry" in df_source.columns:
                    df_export["Name_Clean"] = df_export["Name"].astype(str).str.strip().str.lower()
                    df_source["Name_Clean"] = df_source["Name"].astype(str).str.strip().str.lower()
                    industry_map = df_source.set_index("Name_Clean")["Industry"].to_dict()
                    group_map    = (
                        df_source.set_index("Name_Clean")["Industry Group"].to_dict()
                        if "Industry Group" in df_source.columns else {}
                    )
                    df_export["Industry"]       = df_export["Name_Clean"].map(industry_map).fillna(df_export["Industry"])
                    df_export["Industry Group"] = df_export["Name_Clean"].map(group_map).fillna(df_export["Industry Group"])
                    try:
                        if "Exception_Industry" in df_source.columns:
                            exception_list = df_source["Exception_Industry"].dropna().unique().tolist()
                        else:
                            exception_list = df_source.iloc[:, 29].dropna().unique().tolist()
                        df_export["Exception_Industry"] = df_export["Industry"].apply(
                            lambda x: "Yes" if x in exception_list else "No"
                        )
                    except:
                        df_export["Exception_Industry"] = "No"
                    df_export.drop(columns=["Name_Clean"], inplace=True)
                    print("✅ Enrichment Complete.")
        except Exception as e:
            print(f"⚠️ VLOOKUP Failed: {e}")
            df_export["Exception_Industry"] = "No"

        # ── 6. FILTER LOGIC ──────────────────────────────────────────────
        if "Exception_Industry" not in df_export.columns:
            df_export["Exception_Industry"] = "No"

        condition_op       = df_export["OP Profit > 70%"]               == True
        condition_max      = df_export["Current Profit > in past 3 yrs"] == True
        condition_exc      = df_export["Exception_Industry"]             == "Yes"
        condition_promoter = df_export["Promoter Holding"].fillna(0)    >= 50.0
        condition_pledged  = df_export["Pledged %"].fillna(0)           <= 10.0
        condition_wc       = df_export["WC Days"].fillna(0)             <= 150.0

        has_5_qtrs = "Empty" not in final_quarters[4] if len(final_quarters) >= 5 else False

        if has_5_qtrs:
            mask = (
                condition_op
                & (condition_max | condition_exc)
                & condition_promoter
                & condition_pledged
                & condition_wc
            )
            df_master_candidates = df_export[mask].copy()
        else:
            df_master_candidates = pd.DataFrame()

        # ── 7. PREPARE MASTER APPEND ─────────────────────────────────────
        cols_for_master = df_export.columns[:30].tolist()
        df_to_append    = df_master_candidates[cols_for_master].copy()

        max_sr_no     = 0
        existing_keys = set()
        try:
            if os.path.exists(MASTER_FILE_PATH):
                try:
                    df_current_qtr = pd.read_excel(MASTER_FILE_PATH, sheet_name=latest_quarter_label)
                    if "Sr no." in df_current_qtr.columns:
                        max_sr_no = pd.to_numeric(df_current_qtr["Sr no."], errors='coerce').max()
                        if pd.isna(max_sr_no):
                            max_sr_no = 0
                    if "Company (Page)" in df_current_qtr.columns:
                        existing_keys = set(df_current_qtr["Company (Page)"].astype(str).tolist())
                except Exception:
                    pass
        except Exception:
            pass

        if existing_keys and "Company (Page)" in df_to_append.columns:
            df_to_append = df_to_append[
                ~df_to_append["Company (Page)"].astype(str).isin(existing_keys)
            ].copy()

        if not df_to_append.empty:
            start_seq = int(max_sr_no) + 1
            df_to_append.insert(0, "Sr no.", range(start_seq, start_seq + len(df_to_append)))

        # ── 8. TYPE CONVERSION ───────────────────────────────────────────
        df_to_append = clean_and_convert(df_to_append)
        df_export    = clean_and_convert(df_export)

        # ── 9A. SAVE TO MASTER FILE ──────────────────────────────────────
        try:
            master_exists = os.path.exists(MASTER_FILE_PATH)
            if master_exists:
                ew_kwargs = dict(engine="openpyxl", mode="a", if_sheet_exists="overlay")
            else:
                ew_kwargs = dict(engine="openpyxl", mode="w")

            with pd.ExcelWriter(MASTER_FILE_PATH, **ew_kwargs) as writer:
                if not df_to_append.empty:
                    try:
                        try:
                            ws_qtr    = writer.book[latest_quarter_label]
                            start_row = ws_qtr.max_row
                        except KeyError:
                            start_row = 0

                        df_to_append.to_excel(writer, sheet_name=latest_quarter_label,
                                              index=False, header=(start_row == 0),
                                              startrow=start_row)
                        print(f"✅ Appended {len(df_to_append)} rows to '{latest_quarter_label}'.")

                        ws_qtr        = writer.sheets[latest_quarter_label]
                        row_start_idx = 1 if start_row == 0 else start_row + 1
                        for row in ws_qtr.iter_rows(min_row=row_start_idx, max_row=ws_qtr.max_row):
                            for cell in row:
                                cell.border = THIN_BORDER
                                if cell.col_idx == 1:
                                    cell.alignment = CENTER_ALIGN
                    except Exception as e:
                        print(f"⚠️ Error writing to quarter sheet: {e}")
                        traceback.print_exc()
                else:
                    print(f"⚠️ No new unique candidates for '{latest_quarter_label}'.")

            print(f"✅ Master File saved: {MASTER_FILE_PATH}")
        except PermissionError:
            print(f"\n❌ ERROR: Master File is OPEN. Please close it.\n   {MASTER_FILE_PATH}")

        # ── 9B. SAVE TO RAW DATA FILE (INCREMENTAL QUARTERLY APPEND) ─────
        try:
            raw_file_exists = os.path.exists(RAW_FILE_PATH)
            existing_raw_keys = set()

            # 1. Check if the current quarter's sheet already exists in Raw Data
            if raw_file_exists:
                try:
                    df_existing_raw = pd.read_excel(RAW_FILE_PATH, sheet_name=latest_quarter_label)
                    if "Company (Page)" in df_existing_raw.columns:
                        existing_raw_keys = set(df_existing_raw["Company (Page)"].astype(str).tolist())
                except Exception:
                    pass # Sheet doesn't exist yet

            # 2. Filter the scraped data to ONLY include new companies
            df_raw_to_append = df_export.copy()
            if existing_raw_keys and "Company (Page)" in df_raw_to_append.columns:
                df_raw_to_append = df_raw_to_append[
                    ~df_raw_to_append["Company (Page)"].astype(str).isin(existing_raw_keys)
                ].copy()

            # 3. Append only the new data to the quarter's sheet
            if not df_raw_to_append.empty:
                ew_kwargs = dict(engine="openpyxl", mode="a", if_sheet_exists="overlay") if raw_file_exists else dict(engine="openpyxl", mode="w")
                
                with pd.ExcelWriter(RAW_FILE_PATH, **ew_kwargs) as writer:
                    try:
                        try:
                            ws_raw = writer.book[latest_quarter_label]
                            start_row = ws_raw.max_row
                        except KeyError:
                            start_row = 0

                        df_raw_to_append.to_excel(writer, sheet_name=latest_quarter_label,
                                                  index=False, header=(start_row == 0),
                                                  startrow=start_row)
                        
                        ws_raw = writer.sheets[latest_quarter_label]
                        format_worksheet(ws_raw, header_row=1)
                        
                        print(f"✅ Appended {len(df_raw_to_append)} new companies to Raw Data → '{latest_quarter_label}'")
                    except Exception as e:
                        print(f"⚠️ Error writing to Raw Data sheet: {e}")
                        traceback.print_exc()
            else:
                print(f"✅ Raw Data is up to date. No new companies to add for '{latest_quarter_label}'.")

        except PermissionError:
            print(f"\n❌ ERROR: Raw Data file is OPEN. Please close it.\n   {RAW_FILE_PATH}")


        # ── 10. SAVE RESULTS JSON (for dashboard) ───────────────────────
        try:
            rows = []
            if not df_to_append.empty:
                df_json = df_to_append.copy()
            elif not df_master_candidates.empty:
                df_json = df_master_candidates.copy()
                df_json.insert(0, "Sr no.", range(1, len(df_json) + 1))
            else:
                df_json = pd.DataFrame()

            if not df_json.empty:
                q_cols = [c for c in df_json.columns if any(
                    m in str(c) for m in ["Jan","Feb","Mar","Apr","May","Jun",
                                          "Jul","Aug","Sep","Oct","Nov","Dec"]
                ) and any(ch.isdigit() for ch in str(c)) and ">" not in str(c)]

                latest_q_col = q_cols[0] if len(q_cols) > 0 else None
                prev_q_col   = q_cols[4] if len(q_cols) > 4 else None

                for _, row in df_json.iterrows():
                    rows.append({
                        "sr"           : int(row.get("Sr no.", 0)),
                        "name"         : str(row.get("Name", "")),
                        "company_page" : str(row.get("Company (Page)", "")),
                        "latest_q"     : float(row[latest_q_col]) if latest_q_col and pd.notna(row.get(latest_q_col)) else 0,
                        "prev_q"       : float(row[prev_q_col])   if prev_q_col   and pd.notna(row.get(prev_q_col))   else 0,
                        "op_profit_70" : bool(row.get("OP Profit > 70%", False)),
                        "all_time_hi"  : bool(row.get("Current Profit > in past 3 yrs", False)),
                        "promoter"     : float(row.get("Promoter Holding", 0) or 0),
                        "pledged"      : float(row.get("Pledged %", 0) or 0),
                        "wc_days"      : float(row.get("WC Days", 0) or 0),
                        "industry_group": str(row.get("Industry Group", "") or ""),
                        "industry"     : str(row.get("Industry", "") or ""),
                    })

            import json as _json
            results_payload = _json.dumps({
                "quarter"      : latest_quarter_label,
                "generated_at" : datetime.now().isoformat(),
                "total_count"  : len(rows),
                "rows"         : rows
            }, indent=2)

            results_path = os.path.join(FOLDER_PATH, "results.json")
            with open(results_path, "w") as f:
                f.write(results_payload)
            print(f"✅ results.json saved → {len(rows)} rows")

        except Exception as e:
            print(f"⚠️ results.json generation failed: {e}")
            traceback.print_exc()

        # ── 11. UPLOAD TO CLOUDFLARE R2 ──────────────────────────────────
        print("\n☁️  Uploading files to Cloudflare R2...")
        upload_to_google_drive()

        try:
            results_path = os.path.join(FOLDER_PATH, "results.json")
            if os.path.exists(results_path):
                access_key   = os.environ.get("R2_ACCESS_KEY_ID", "")
                secret_key   = os.environ.get("R2_SECRET_ACCESS_KEY", "")
                endpoint_url = os.environ.get("R2_ENDPOINT_URL", "")
                bucket_name  = os.environ.get("R2_BUCKET_NAME", "quarterly-profit-surge")

                if all([access_key, secret_key, endpoint_url]):
                    import boto3
                    from botocore.config import Config
                    s3 = boto3.client(
                        "s3",
                        endpoint_url          = endpoint_url,
                        aws_access_key_id     = access_key,
                        aws_secret_access_key = secret_key,
                        config                = Config(signature_version="s3v4"),
                        region_name           = "auto"
                    )
                    s3.upload_file(
                        results_path, bucket_name, "results.json",
                        ExtraArgs={"ContentType": "application/json"}
                    )
                    print("✅ Uploaded results.json to Cloudflare R2.")
        except Exception as e:
            print(f"⚠️ results.json R2 upload failed: {e}")
            traceback.print_exc()

finally:
    try:
        driver.quit()
    except:
        pass